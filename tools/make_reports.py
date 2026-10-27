"""Create reports from the retained title-page template and task Markdown."""
from pathlib import Path
from copy import deepcopy
import re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = Path(r'C:\Users\artem\Downloads\laba_2_7.docx')

def replace_text(p, text):
    props = deepcopy(p.runs[-1]._r.rPr) if p.runs and p.runs[-1]._r.rPr is not None else None
    p.clear()
    r = p.add_run(text)
    if props is not None:
        r._r.insert(0, props)
    return r

def report(topic):
    d = Document(TEMPLATE)
    d.core_properties.author='Дубограй Артем Юрьевич'
    d.core_properties.last_modified_by='Дубограй Артем Юрьевич'
    d.core_properties.title=topic
    d.core_properties.subject='Информатика Вариант 6 Группа 615'
    # The first 51 paragraphs are the supplied cover. Remove sample solutions.
    for p in list(d.paragraphs)[51:]:
        p._element.getparent().remove(p._element)
    for t in list(d.tables):
        t._element.getparent().remove(t._element)
    ps = d.paragraphs
    r = replace_text(ps[17], 'Лабораторная работа по теме «' + topic + '»\nпо курсу «Информатика»')
    r.font.size = Pt(14)
    r.italic=False
    replace_text(ps[30], '\tстудент гр. 615\n\tДубограй Артем Юрьевич')
    ps[30].alignment=WD_ALIGN_PARAGRAPH.LEFT
    replace_text(ps[33], '\tНикичкин Б. В.')
    # Compensate for the additional author line within the existing cover grid.
    for i in [49,48]:ps[i]._element.getparent().remove(ps[i]._element)
    st = d.styles['Normal']
    for name in ['Heading 1','Heading 2','Title']:
        s=d.styles[name] if name in d.styles else d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
        s.base_style=st;s.font.size=Pt(14);s.font.bold=True
        if name.startswith('Heading'):
            pp=s.element.get_or_add_pPr();ol=OxmlElement('w:outlineLvl');ol.set(qn('w:val'),'0' if name=='Heading 1' else '1');pp.append(ol)
    return d

def para(d, text='', style=None):
    p=d.add_paragraph(style=style)
    p.paragraph_format.space_after=Pt(6)
    p.paragraph_format.line_spacing=1.15
    p.paragraph_format.keep_together=False
    # Bold spans, other Markdown is intentionally not interpreted.
    for i,s in enumerate(re.split(r'\*\*(.*?)\*\*',text)):
        r=p.add_run(s);r.bold=bool(i%2) or bool(style);r.font.name='Times New Roman';r.font.size=Pt(14 if style else 12)
    return p

def table(d, rows):
    t=d.add_table(rows=0,cols=len(rows[0]));t.autofit=True
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for c,text in zip(cells,row):
            c.text=text
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10);r.bold=(ri==0)
            pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
            pr.append(b)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EAEAEA');pr.append(sh)
        if ri==0:
            trpr=t.rows[0]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:tblHeader'))
    return t

def from_markdown(path,topic):
    d=report(topic)
    first=True;code=False;rows=[];page=False
    for line in path.read_text(encoding='utf-8').splitlines():
        if line.startswith('# '):continue
        if line.startswith('|') and not code:
            rr=[x.strip() for x in line.strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+',x) for x in rr):rows.append(rr)
            continue
        if rows:table(d,rows);rows=[]
        if line.startswith('```'):code=not code;continue
        if line=='<!-- PAGE -->':page=True;continue
        if line.startswith('!['):
            target=re.search(r'\]\((.*?)\)',line).group(1)
            p=d.add_paragraph();p.alignment=1;p.add_run().add_picture(str(path.parent/target),width=Cm(15.5))
            continue
        if not line.strip():continue
        if line.startswith('## '):
            p=para(d,line[3:],'Heading 1')
            p.paragraph_format.keep_with_next=True
        elif line.startswith('### '):
            p=para(d,line[4:],'Heading 2');p.paragraph_format.keep_with_next=True
        else:p=para(d,line)
        if first:p.paragraph_format.page_break_before=True;first=False
        if page:p.paragraph_format.page_break_before=True;page=False
        if code:
            p.paragraph_format.line_spacing=1
            p.paragraph_format.space_after=Pt(0)
            for r in p.runs:r.font.name='Consolas';r.font.size=Pt(10)
    if rows:table(d,rows)
    out=path.with_name('Дубограй_А_Ю_Отчет.docx');d.save(out)
    print(out)
    return out

if __name__=='__main__':
    import sys
    sys.stdout.reconfigure(encoding='utf-8')
    for arg in sys.argv[1:]:
        p=Path(arg)
        topic=p.read_text(encoding='utf-8').splitlines()[0].removeprefix('# ')
        from_markdown(p,topic)
