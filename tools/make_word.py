from pathlib import Path
import sys
sys.stdout.reconfigure(encoding='utf-8')
from copy import deepcopy
from docx.shared import Cm,Pt,RGBColor
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_UNDERLINE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader
from make_reports import report,para,ROOT

OUT=ROOT/'05_Word';OUT.mkdir(exist_ok=True)
d=report('Создание и форматирование текстового документа')
s=d.add_section(WD_SECTION_START.NEW_PAGE)
s.top_margin=s.bottom_margin=s.left_margin=Cm(2);s.right_margin=Cm(1)
s.header.is_linked_to_previous=False
p=s.header.paragraphs[0];p.alignment=2;p.add_run('-')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld);p.add_run('-')

def heading(n,title):
 p=para(d,f'Задание {n} {title}','Heading 1');p.paragraph_format.page_break_before=True;p.paragraph_format.keep_with_next=True
 if n==1:p.paragraph_format.page_break_before=False
 return p

def mr(text):
 r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=text;r.append(t);return r
def box(tag,children):
 e=OxmlElement('m:'+tag)
 if not isinstance(children,list):children=[children]
 for c in children:e.append(c)
 return e
def frac(a,b):return box('f',[box('num',a),box('den',b)])
def bar(children):
 pr=OxmlElement('m:barPr');pos=OxmlElement('m:pos');pos.set(qn('m:val'),'top');pr.append(pos)
 return box('bar',[pr,box('e',children)])
def root(children):
 pr=OxmlElement('m:radPr');hide=OxmlElement('m:degHide');hide.set(qn('m:val'),'1');pr.append(hide)
 return box('rad',[pr,OxmlElement('m:deg'),box('e',children)])
def mathline(children):
 p=d.add_paragraph();p.paragraph_format.space_after=Pt(18)
 m=box('oMath',children);p._p.append(m);return p

heading(1,'Форматирование абзацев')
para(d,'Вариант 6. Наборы оформления для четырёх абзацев: 6, 5, 4, 3.')
text=PdfReader(r'F:\study\методичка\Лаб_раб по Word.pdf').pages[0].extract_text()
start=text.index('..На краю');end=text.index('Оформите первый')
raw=text[start:end]
begins=['..На краю','Этот дуб','Князь Андрей','«Да, он прав']
ix=[raw.index(k) for k in begins]+[len(raw)]
settings=[(2,4,0,1,.5,1,2),(3,0,-1,1,0,3,1.5),(3,3,1.27,0,1,2,1.3),(6,0,0,.63,1,3,1)]
for i in range(4):
 t=' '.join(raw[ix[i]:ix[i+1]].split()).lstrip('.')
 p=para(d,t);b,a,first,left,right,align,line=settings[i]
 pf=p.paragraph_format;pf.space_before=Pt(b);pf.space_after=Pt(a);pf.first_line_indent=Cm(first);pf.left_indent=Cm(left);pf.right_indent=Cm(right);pf.line_spacing=line;p.alignment=align
 for r in p.runs:r.font.size=Pt(12)

heading(2,'Математические формулы')
para(d,'Вариант 6: формулы 2, 6, 7 и 10. Все формулы представлены редактируемыми математическими объектами Word.')
para(d,'Формула 2')
rows=[
 [mr('tg'),frac(mr('x+y'),mr('1−x·y')),mr(',   при xy<1')],
 [mr('−3.14,   при xy>1, x≥0')],
 [mr('1+sin(x+y),   при x<0, xy>1')]
]
em=box('eqArr',[box('e',row) for row in rows])
dpr=OxmlElement('m:dPr');ch=OxmlElement('m:begChr');ch.set(qn('m:val'),'{');dpr.append(ch);ch=OxmlElement('m:endChr');ch.set(qn('m:val'),'');dpr.append(ch)
mathline([mr('z='),box('d',[dpr,box('e',em)])])
para(d,'Формула 6')
npr=OxmlElement('m:naryPr');ch=OxmlElement('m:chr');ch.set(qn('m:val'),'∫');npr.append(ch);ch=OxmlElement('m:limLoc');ch.set(qn('m:val'),'subSup');npr.append(ch)
mathline([box('nary',[npr,box('sub',mr('4')),box('sup',mr('8')),box('e',root(mr('x+5')))]),mr(' dx')])
para(d,'Формула 7')
cf=frac(mr('1'),mr('11'))
for n in range(10,0,-1):cf=frac(mr('1'),[mr(str(n)+'+'),cf])
mathline([mr('y='),cf])
para(d,'Формула 10')
mathline([bar([mr('(X→Y)·'),bar([bar(mr('X')),mr('Z')]),mr('·((X→'),bar(mr('Y')),mr(')+'),bar(mr('Z')),mr(')')])])

heading(3,'Специальные символы')
para(d,'Набор 6. Размеры символов: 20, 22, 24, 26, 28 и 30 пт; каждый символ имеет отдельный цвет.')
p=d.add_paragraph()
for i,c in enumerate(['\uf024','\uf0c9','\uf07a','\uf044','\uf063','§']):
 r=p.add_run(c+'  ');r.font.name='Wingdings' if i<5 else 'Times New Roman';r.font.size=Pt(20+2*i);r.font.color.rgb=RGBColor.from_string(['174E85','B52D3A','1E7043','784F9C','AD6600','663333'][i])

heading(4,'Объект WordArt')
para(d,'Текст надписи: Дубограй Артем. Форма — дуга вверх, заливка «Пламя», жёлтая сплошная граница толщиной 2,25 пт.')
para(d,'[[WORDART]]')
para(d,'Рисунок 4 Декоративная надпись WordArt')

heading(5,'Члены предложения')
words=[('Черные',WD_UNDERLINE.WAVY),('вороны',WD_UNDERLINE.SINGLE),('вышагивают',WD_UNDERLINE.DOUBLE),('рядами,',WD_UNDERLINE.DOT_DASH),('как на параде,',WD_UNDERLINE.DOT_DASH),('и',None),('червяков',WD_UNDERLINE.DASH),('склевывают.',WD_UNDERLINE.DOUBLE)]
p=d.add_paragraph()
for t,u in words:
 r=p.add_run(t);r.underline=u;r.font.name='Times New Roman';r.font.size=Pt(14);p.add_run(' ')
para(d,'Черные — определение; вороны — подлежащее; вышагивают, склевывают — однородные сказуемые; рядами и как на параде — обстоятельства; червяков — дополнение; и — союз.')

heading(6,'Принципиальная схема из автофигур')
para(d,'Вариант 6. Схема с фоторезистором LDR, потенциометром R1, делителем R2–R3, усилителем U1 и выходным резистором R4. Редактируемые автофигуры объединены на одном полотне рисунка, которое перемещается как единый объект.')
para(d,'[[CIRCUIT]]')
para(d,'Рисунок 6 Принципиальная схема варианта 6')

heading(7,'Форматирование таблицы')
para(d,'Таблица 7 Учебные значения энергии и пищевых веществ по образцу методички')
table=d.add_table(rows=3,cols=10);table.autofit=False
for c in table.columns:c.width=Cm(1.8)
table.cell(0,0).merge(table.cell(2,0)).text='Группа труда';table.cell(0,1).merge(table.cell(2,1)).text='Возраст, годы'
table.cell(0,2).merge(table.cell(0,5)).text='Мужчины';table.cell(0,6).merge(table.cell(0,9)).text='Женщины'
for j in range(2,10):table.cell(1,j).merge(table.cell(2,j)).text=['Энергия, ккал','Белки, г','Жиры, г','Углеводы, г'][(j-2)%4]
data=[
 [2450,72,81,358,2000,61,67,289],[2300,68,77,335,1900,59,63,274],[2100,65,70,303,1800,58,60,257],
 [2800,80,93,411,2200,66,73,318],[2650,77,88,387,2150,65,72,311],[2500,72,83,366,2100,63,70,305],
 [3300,94,110,484,2600,76,87,378],[3150,89,105,462,2550,74,85,372],[2950,84,98,432,2500,72,83,366],
 [3850,108,128,566,3050,87,102,452],[3600,102,120,528,2950,84,98,432],[3400,96,113,499,2850,82,95,417],
 [4200,117,154,586,'—','—','—','—'],[3950,110,144,550,'—','—','—','—'],[3750,104,137,524,'—','—','—','—']]
for i,row in enumerate(data):
 cells=table.add_row().cells
 for c,t in zip(cells,[i//3+1,['18–19','30–39','40–59'][i%3]]+row):c.text=str(t)
for g in range(5):table.cell(3+g*3,0).merge(table.cell(5+g*3,0)).text=str(g+1)
colors=['CCFFFF','CCFFFF','CCFFFF','CCFFCC','CCFFCC','CCFFCC','CCFFFF','FFFFCC','CCFFFF','FFFFCC']
for ri,row in enumerate(table.rows):
 for j,c in enumerate(row.cells):
  pr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'FFFFCC' if ri<3 else colors[j]);pr.append(sh)
  bd=OxmlElement('w:tcBorders')
  for e in ['top','bottom','left','right']:
   b=OxmlElement('w:'+e);outside=(e=='top' and ri==0)or(e=='bottom'and ri==17)or(e=='left'and j==0)or(e=='right'and j==9)
   b.set(qn('w:val'),'thinThickSmallGap' if outside else 'dotDash');b.set(qn('w:color'),'333F50' if outside else '000080');b.set(qn('w:sz'),'12' if outside else '6');bd.append(b)
  pr.append(bd)
  c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP if j<2 else WD_CELL_VERTICAL_ALIGNMENT.BOTTOM if 5<=j<=8 else WD_CELL_VERTICAL_ALIGNMENT.CENTER
  for p in c.paragraphs:
   p.alignment=0 if 2<=j<=4 else 1
   p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
   for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10);r.bold=ri<3
 for r in table.rows[:3]:
  tr=r._tr.get_or_add_trPr()
  if tr.find(qn('w:tblHeader')) is None:tr.append(OxmlElement('w:tblHeader'))
para(d,'Оформление варианта 6: внешняя линия — тёмно-сизая, тип 5; внутренняя — тёмно-синяя, тип 3. Шапка — заливка 7. Выравнивание столбцов: 1–2 — набор 2; 3–5 — набор 4; 6–9 — набор 8; 10 — набор 5.')
para(d,'Вывод: выполнены задания по форматированию абзацев, вводу формул и символов, созданию WordArt и сгруппированной схемы, оформлению таблицы. Применены единые стили заголовков, параметры страницы варианта 6 и автоматическое оглавление.')

p=para(d,'Оглавление');p.runs[0].bold=True;p.runs[0].font.size=Pt(14);p.paragraph_format.page_break_before=True
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z \\u');p=d.add_paragraph();p._p.append(fld)
out=OUT/'Дубограй_А_Ю_Word.docx';d.save(out)
print(out)
